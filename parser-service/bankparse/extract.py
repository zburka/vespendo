"""Fájlból nyers sorok kinyerése: CSV, XLSX, PDF (táblázatos és szöveges)."""
import csv
import io
import re
from datetime import date, datetime

from . import cells


class ParseHiba(Exception):
    """Érthető, felhasználónak megjeleníthető hiba."""


Sorok = list[list[object]]


def kinyer(adat: bytes, fajlnev: str) -> Sorok:
    nev = (fajlnev or "").lower()
    if adat[:4] == b"%PDF" or nev.endswith(".pdf"):
        return _pdf(adat)
    if adat[:2] == b"PK" or nev.endswith(".xlsx"):
        return _xlsx(adat)
    if nev.endswith(".xls"):
        raise ParseHiba("A régi .xls formátum nem támogatott, mentsd el .xlsx vagy .csv fájlként.")
    return _csv(adat)


def _dekodol(adat: bytes) -> str:
    try:
        return adat.decode("utf-8-sig")
    except UnicodeDecodeError:
        return adat.decode("cp1250", errors="replace")


def _csv(adat: bytes) -> Sorok:
    szoveg = _dekodol(adat)
    sorok = [s for s in szoveg.splitlines() if s.strip()]
    if not sorok:
        raise ParseHiba("A fájl üres.")
    minta = sorok[:50]
    elvalaszto = max(";,\t|", key=lambda d: sum(s.count(d) for s in minta))
    return [list(r) for r in csv.reader(io.StringIO("\n".join(sorok)), delimiter=elvalaszto)]


def _xlsx(adat: bytes) -> Sorok:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(adat), read_only=True, data_only=True)
    except Exception as e:  # sérült vagy jelszavas fájl
        raise ParseHiba("Az XLSX fájl nem olvasható (sérült vagy jelszóval védett).") from e
    sorok: Sorok = []
    for lap in wb.worksheets:
        for sor in lap.iter_rows(values_only=True):
            if any(not cells.ures(c) for c in sor):
                sorok.append(["" if c is None else c for c in sor])
    return sorok


# --- PDF -------------------------------------------------------------------------------------

# erős összegminta: előjeles, ezres tagolású vagy tizedes
_OSSZEG = r"[+-]?\d{1,3}(?:[  .]\d{3})+(?:,\d{1,2})?|[+-]?\d+[.,]\d{1,2}|[+-]\d+"
_PENZNEM = r"(?:\s*(?:HUF|Ft|EUR|USD))?"
_VEGE = re.compile(rf"((?:\s+(?:{_OSSZEG}){_PENZNEM})+)\s*$", re.IGNORECASE)
_SZAMOK = re.compile(_OSSZEG)


def _pdf(adat: bytes) -> Sorok:
    import pdfplumber

    try:
        pdf = pdfplumber.open(io.BytesIO(adat))
    except Exception as e:
        raise ParseHiba("A PDF nem olvasható (sérült vagy jelszóval védett).") from e

    tablazat: Sorok = []
    oldalak_sorai: list[list[str]] = []
    szoveg_van = False
    with pdf:
        for oldal in pdf.pages:
            for t in oldal.extract_tables():
                for sor in t:
                    tisztitott = [(c or "").replace("\n", " ").strip() for c in sor]
                    if any(tisztitott):
                        tablazat.append(tisztitott)
            szoveg = oldal.extract_text() or ""
            szoveg_van = szoveg_van or bool(szoveg.strip())
            oldalak_sorai.append(_sorok_szavakbol(oldal))

    if not szoveg_van:
        raise ParseHiba("A PDF képként tartalmaz szöveget (szkennelt), ez nem támogatott.")
    # táblázatos mód, ha a táblák tartalmaznak dátumot; egyébként soronkénti szövegelemzés
    if sum(1 for r in tablazat if any(cells.datum(c) for c in r)) >= 3:
        return tablazat
    return _szoveges_sorok(oldalak_sorai)


_MOJIBAKE_JELEK = "ß÷Ř"


def _mojibake_javit(szoveg: str) -> str:
    """Egyes banki PDF-ek betűkódolása elromlott (cp1250 bájtok cp852-ként olvasva):
    'Tßp-termÚny' -> 'Táp-termény'. Csak a félreérthetetlen jeleket tartalmazó szavakat javítja."""
    def javit(m: re.Match) -> str:
        szo = m.group(0)
        if not any(j in szo for j in _MOJIBAKE_JELEK):
            return szo
        try:
            return szo.encode("cp852").decode("cp1250")
        except (UnicodeEncodeError, UnicodeDecodeError):
            return szo
    return re.sub(r"\S+", javit, szoveg)


def _szamok_a_sor_vegen(maradek: str) -> tuple[str, list[str]]:
    """Az oszlophatárral (tabulátor) elválasztott záró számmezők; ennek híján erős összegminták."""
    mezok = [m for m in maradek.split("\t")]
    szamok: list[str] = []
    while len(mezok) > 1:
        mezo = mezok[-1].strip()
        if mezo and cells.osszeg(mezo) is not None and cells.datum(mezo) is None:
            szamok.insert(0, mezo)
            mezok.pop()
        else:
            break
    if szamok:
        return "\t".join(mezok), szamok
    v = _VEGE.search(maradek)
    if v:
        return maradek[: v.start()], [x.strip() for x in _SZAMOK.findall(v.group(1))]
    return maradek, []


def _szoveges_sorok(oldalak: list[list[str]]) -> Sorok:
    """[dátum, leírás, szám1, szám2, ...] sorok. A dátum nélküli sorok az előző tétel leírásához
    fűződnek, de csak az oldalon belül: az oldal fejléce az első dátumos sor előtt eldobódik."""
    eredmeny: Sorok = []
    for oldal in oldalak:
        utolso: list[object] | None = None
        for nyers in oldal:
            sor = nyers.replace("\u00a0", " ")
            m = cells.DATUM_SOR_ELEJE.match(sor)
            if not m:
                if utolso is not None and sor.strip() and not _VEGE.search(sor):
                    utolso[1] = f"{utolso[1]} {sor.strip().replace(chr(9), ' ')}".strip()
                continue
            maradek = sor[m.end():]
            masodik = cells.DATUM_SOR_ELEJE.match(maradek)  # értéknap: a könyvelési dátum marad
            if masodik:
                maradek = maradek[masodik.end():]
            leiras, szamok = _szamok_a_sor_vegen(maradek)
            leiras = _mojibake_javit(re.sub(r"\s+", " ", leiras.replace("\t", " ")).strip())
            utolso = [m.group(1), leiras, *szamok]
            eredmeny.append(utolso)
    return eredmeny


def _sorok_szavakbol(oldal) -> list[str]:
    """Szavak pozíciójából sorok; a nagy vízszintes hézag tabulátor (oszlophatár), így az
    egymás melletti számok ('20 000    480 000') nem olvadnak össze."""
    szavak = oldal.extract_words(extra_attrs=["size"], keep_blank_chars=False)
    sorok: list[list[dict]] = []
    for sz in sorted(szavak, key=lambda w: (round(w["top"]), w["x0"])):
        if sorok and abs(sorok[-1][0]["top"] - sz["top"]) <= 3:
            sorok[-1].append(sz)
        else:
            sorok.append([sz])
    kimenet = []
    for sor in sorok:
        sor.sort(key=lambda w: w["x0"])
        szoveg = sor[0]["text"]
        for elozo, w in zip(sor, sor[1:]):
            hezag = w["x0"] - elozo["x1"]
            szoveg += ("\t" if hezag > 0.45 * w["size"] else " ") + w["text"]
        kimenet.append(szoveg)
    return kimenet
